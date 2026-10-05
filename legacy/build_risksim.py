"""Genera MiRiskSim.xlsx: simulador Monte Carlo 100% con fórmulas (Excel Mac/Win, LibreOffice, WPS)."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter as L

N = 1000          # simulaciones
R0 = 5            # primera fila de datos en Simulacion
R1 = R0 + N - 1

F = "Arial"
f_norm = Font(name=F, size=10)
f_bold = Font(name=F, size=10, bold=True)
f_title = Font(name=F, size=16, bold=True, color="1F3864")
f_sub = Font(name=F, size=10, italic=True, color="595959")
f_in = Font(name=F, size=10, color="0000FF")
f_head = Font(name=F, size=10, bold=True, color="FFFFFF")
f_link = Font(name=F, size=10, color="008000")
fill_head = PatternFill("solid", fgColor="1F3864")
fill_in = PatternFill("solid", fgColor="FFF2CC")
fill_sec = PatternFill("solid", fgColor="D9E2F3")
thin = Side(style="thin", color="BFBFBF")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
wrap = Alignment(wrap_text=True, vertical="top")

wb = Workbook()

def style_range(ws, rng, font=None, fill=None, border=None, align=None, fmt=None):
    for row in ws[rng]:
        for c in row:
            if font: c.font = font
            if fill: c.fill = fill
            if border: c.border = border
            if align: c.alignment = align
            if fmt: c.number_format = fmt

def header(ws, row, c0, labels):
    for i, t in enumerate(labels):
        c = ws.cell(row=row, column=c0 + i, value=t)
        c.font, c.fill, c.alignment, c.border = f_head, fill_head, center, box

# ------------------------------------------------------------------ Instrucciones
ws = wb.active
ws.title = "Instrucciones"
ws.sheet_view.showGridLines = False
ws["A1"] = "MiRiskSim – Simulación Monte Carlo de un proyecto"
ws["A1"].font = f_title
lines = [
    ("Qué es", "Simulador de riesgo para evaluación de proyectos hecho solo con fórmulas (sin macros). Funciona en Excel (Mac/Windows), LibreOffice y WPS."),
    ("Cómo se usa", "1) Ve a la hoja 'Supuestos' y edita las celdas azules/amarillas: distribución y parámetros de cada variable incierta.\n"
                    "2) Ve a 'Resultados': ahí aparecen estadísticos, probabilidades, histograma y sensibilidad.\n"
                    "3) Para volver a simular pulsa F9 (Excel Mac: fn+F9 o Cmd+=; LibreOffice: Ctrl+Shift+F9). Cada recálculo = 1000 nuevas corridas."),
    ("Modelo incluido", "Flujo de caja a 5 años:  Ingresos_t = Ingresos_1 × (1+g)^(t-1);  Flujo_t = Ingresos_t × (1 − Costos%) × (1 − Impuesto);  año 5 suma el valor de rescate.\n"
                        "VAN = −Inversión + Σ Flujo_t /(1+r)^t.   TIR calculada por cada corrida."),
    ("Distribuciones", "Normal (P1=media, P2=desv. estándar) · Triangular (P1=mínimo, P2=más probable, P3=máximo) · Uniforme (P1=mínimo, P2=máximo) · "
                       "Lognormal (P1=media, P2=desv. estándar, ambas de la variable) · Fija (P1=valor)."),
    ("Truncar", "Las columnas Mín/Máx recortan los valores simulados (por ejemplo, una tasa no puede ser negativa). Déjalas vacías si no quieres límites."),
    ("Congelar resultados", "RAND() cambia en cada recálculo. Para guardar una corrida para tu informe: selecciona Simulacion!I5:W1004, copia y pega como VALORES, o toma capturas."),
    ("Otro modelo", "Para tu propio proyecto cambia las fórmulas de las columnas P:W de la hoja 'Simulacion' (flujos, VAN, TIR) y agrega/quita variables en 'Supuestos'. "
                    "Más corridas: arrastra hacia abajo la última fila de 'Simulacion' y amplía los rangos en 'Resultados'."),
    ("Limitaciones", "No incluye correlaciones entre variables ni optimización/pronóstico de series de tiempo como Risk Simulator. Las variables se muestrean de forma independiente."),
    ("Colores", "Texto azul = dato que puedes editar · Fondo amarillo = supuesto clave · Texto negro = fórmula."),
]
r = 3
for k, v in lines:
    ws.cell(row=r, column=1, value=k).font = f_bold
    c = ws.cell(row=r, column=2, value=v); c.font = f_norm; c.alignment = wrap
    ws.cell(row=r, column=1).alignment = wrap
    ws.row_dimensions[r].height = max(30, 15 * (v.count("\n") + 1 + len(v) // 110))
    r += 1
ws.column_dimensions["A"].width = 22
ws.column_dimensions["B"].width = 110

# ------------------------------------------------------------------ Supuestos
sp = wb.create_sheet("Supuestos")
sp.sheet_view.showGridLines = False
sp["A1"] = "Supuestos y variables inciertas"; sp["A1"].font = f_title
sp["A2"] = "Edita las celdas azules. Los valores base alimentan el caso determinista; la simulación usa la distribución elegida."
sp["A2"].font = f_sub
header(sp, 4, 1, ["Variable", "Unidad", "Distribución", "P1", "P2", "P3", "Mín (truncar)", "Máx (truncar)", "Valor base", "Significado de P1 / P2 / P3"])

# nombre, unidad, dist, p1,p2,p3, min,max, base, formato
VARS = [
    ("Inversión inicial",          "u.m.", "Triangular", 90000, 100000, 125000, None, None, 100000, '#,##0'),
    ("Ingresos año 1",             "u.m.", "Normal",     65000, 10000,  None,   0,    None, 65000,  '#,##0'),
    ("Crecimiento anual ingresos", "%",    "Uniforme",   0.00,  0.08,   None,   None, None, 0.04,   '0.0%'),
    ("Costos (% de ingresos)",     "%",    "Triangular", 0.38,  0.45,   0.60,   0,    1,    0.45,   '0.0%'),
    ("Tasa de impuesto",           "%",    "Fija",       0.25,  None,   None,   None, None, 0.25,   '0.0%'),
    ("Tasa de descuento",          "%",    "Normal",     0.12,  0.015,  None,   0.01, None, 0.12,   '0.0%'),
    ("Valor de rescate (año 5)",   "u.m.", "Lognormal",  15000, 5000,   None,   0,    None, 15000,  '#,##0'),
]
V0 = 5  # primera fila de variables en Supuestos
meaning = {
    "Triangular": "mín / más probable / máx", "Normal": "media / desv. estándar / –",
    "Uniforme": "mín / máx / –", "Lognormal": "media / desv. estándar / –", "Fija": "valor / – / –",
}
dv = DataValidation(type="list", formula1='"Normal,Triangular,Uniforme,Lognormal,Fija"', allow_blank=False)
sp.add_data_validation(dv)
for i, (n, u, d, p1, p2, p3, mn, mx, base, fmt) in enumerate(VARS):
    rr = V0 + i
    vals = [n, u, d, p1, p2, p3, mn, mx, base]
    for j, v in enumerate(vals):
        c = sp.cell(row=rr, column=1 + j, value=v)
        c.border = box
        c.font = f_bold if j == 0 else (f_norm if j == 1 else f_in)
        if j >= 3:
            c.number_format = fmt
        if j == 2:
            c.alignment = center
    sp.cell(row=rr, column=10, value=f'=IF(C{rr}="Triangular","mín / más probable / máx",IF(C{rr}="Normal","media / desv. estándar",IF(C{rr}="Uniforme","mín / máx",IF(C{rr}="Lognormal","media / desv. estándar","valor"))))').font = f_norm
    sp.cell(row=rr, column=10).border = box
    sp.cell(row=rr, column=9).fill = fill_in
    dv.add(sp.cell(row=rr, column=3))
VL = V0 + len(VARS) - 1   # última fila de variables (11)

sp["A13"] = "Nota: 'u.m.' = unidades monetarias (usa tu moneda). Valores de ejemplo, reemplázalos por los de tu proyecto."
sp["A13"].font = f_sub

# Caso determinista
sp["A15"] = "Caso determinista (con valores base)"; sp["A15"].font = f_bold
style_range(sp, "A15:G15", fill=fill_sec)
header(sp, 16, 1, ["Concepto", "Año 0", "Año 1", "Año 2", "Año 3", "Año 4", "Año 5"])
sp["A17"], sp["A18"], sp["A19"] = "Ingresos", "Flujo neto", "Flujo para TIR/VAN"
for a in ("A17", "A18", "A19"):
    sp[a].font = f_bold; sp[a].border = box
base = lambda i: f"$I${V0 + i}"
for t in range(0, 6):
    col = L(2 + t)
    if t == 0:
        sp[f"{col}17"] = 0
        sp[f"{col}18"] = f"=-{base(0)}"
        sp[f"{col}19"] = f"={col}18"
    else:
        sp[f"{col}17"] = f"={base(1)}*(1+{base(2)})^({t}-1)"
        sp[f"{col}18"] = f"={col}17*(1-{base(3)})*(1-{base(4)})" + (f"+{base(6)}" if t == 5 else "")
        sp[f"{col}19"] = f"={col}18"
    for rw in (17, 18, 19):
        c = sp[f"{col}{rw}"]; c.number_format = '#,##0;(#,##0);-'; c.font = f_norm; c.border = box
sp["A21"], sp["A22"] = "VAN base", "TIR base"
sp["B21"] = f"=B19+NPV({base(5)},C19:G19)"; sp["B21"].number_format = '#,##0;(#,##0);-'
sp["B22"] = "=IFERROR(IRR(B19:G19,0.1),\"n/d\")"; sp["B22"].number_format = '0.0%'
for a in ("A21", "A22"): sp[a].font = f_bold
for a in ("B21", "B22"): sp[a].font = f_bold; sp[a].border = box
for col, w in zip("ABCDEFGHIJ", (30, 9, 14, 12, 12, 12, 14, 14, 12, 30)):
    sp.column_dimensions[col].width = w

# ------------------------------------------------------------------ Simulacion
sm = wb.create_sheet("Simulacion")
sm["A1"] = "Corridas Monte Carlo (cada fila es un escenario completo)"; sm["A1"].font = f_title
sm["A2"] = "Se recalcula con F9. Columnas B:H = números aleatorios · I:O = valores muestreados · P:U = flujos · V = VAN · W = TIR"
sm["A2"].font = f_sub
labels = ["#"] + [f"U {v[0][:14]}" for v in VARS] + [v[0] for v in VARS] + \
         ["Flujo 0", "Flujo 1", "Flujo 2", "Flujo 3", "Flujo 4", "Flujo 5", "VAN", "TIR"]
header(sm, 4, 1, labels)
sm.row_dimensions[4].height = 42
sm.freeze_panes = "B5"

def sample_formula(i, ucell):
    sr = V0 + i
    D, E, Fc, G, H = (f"Supuestos!$D${sr}", f"Supuestos!$E${sr}", f"Supuestos!$F${sr}",
                      f"Supuestos!$G${sr}", f"Supuestos!$H${sr}")
    C = f"Supuestos!$C${sr}"
    s2 = f"LN(1+{E}^2/{D}^2)"
    raw = (f'IF({C}="Normal",NORMINV({ucell},{D},{E}),'
           f'IF({C}="Triangular",IF({ucell}<({E}-{D})/({Fc}-{D}),{D}+SQRT({ucell}*({Fc}-{D})*({E}-{D})),{Fc}-SQRT((1-{ucell})*({Fc}-{D})*({Fc}-{E}))),'
           f'IF({C}="Uniforme",{D}+{ucell}*({E}-{D}),'
           f'IF({C}="Lognormal",LOGINV({ucell},LN({D})-{s2}/2,SQRT({s2})),{D}))))')
    return f'=MIN(IF({H}="",1E+99,{H}),MAX(IF({G}="",-1E+99,{G}),{raw}))'

nv = len(VARS)
for k in range(N):
    r = R0 + k
    sm.cell(row=r, column=1, value=k + 1)
    for i in range(nv):
        ucol = L(2 + i)
        sm[f"{ucol}{r}"] = "=RAND()*(1-2E-6)+1E-6"
        sm[f"{L(2 + nv + i)}{r}"] = sample_formula(i, f"{ucol}{r}")
    # columnas de valores: I=Inv J=Ing1 K=g L=costo M=tax N=r O=rescate
    sm[f"P{r}"] = f"=-I{r}"
    for t in range(1, 6):
        col = L(16 + t)  # Q..U
        extra = f"+$O{r}" if t == 5 else ""
        sm[f"{col}{r}"] = f"=$J{r}*(1+$K{r})^({t}-1)*(1-$L{r})*(1-$M{r}){extra}"
    sm[f"V{r}"] = f"=P{r}+NPV($N{r},Q{r}:U{r})"
    sm[f"W{r}"] = f'=IFERROR(IRR(P{r}:U{r},0.1),"")'
    for c in range(1, 24):
        cell = sm.cell(row=r, column=c)
        cell.font = f_norm
        if 2 <= c <= 8: cell.number_format = '0.000'
        elif 9 <= c <= 15: cell.number_format = VARS[c - 9][9]
        elif 16 <= c <= 22: cell.number_format = '#,##0;(#,##0);-'
        elif c == 23: cell.number_format = '0.0%'
for c in range(1, 24):
    sm.column_dimensions[L(c)].width = 11 if c > 1 else 7
for c in range(2, 9):
    sm.column_dimensions[L(c)].hidden = False
sm.column_dimensions.group("B", "H", hidden=True)  # aleatorios ocultos (agrupados)

# ------------------------------------------------------------------ Resultados
rs = wb.create_sheet("Resultados", 1)
rs.sheet_view.showGridLines = False
rs["A1"] = "Resultados de la simulación"; rs["A1"].font = f_title
rs["A2"] = "Pulsa F9 para generar una nueva simulación. Los valores cambian ligeramente en cada corrida (es normal)."
rs["A2"].font = f_sub
VAN = f"Simulacion!$V${R0}:$V${R1}"
TIR = f"Simulacion!$W${R0}:$W${R1}"

header(rs, 4, 1, ["Estadístico (VAN)", "Valor"])
stats = [
    ("Nº de simulaciones", f"=COUNT({VAN})", '#,##0'),
    ("VAN determinista (caso base)", "=Supuestos!B21", '#,##0;(#,##0);-'),
    ("Media", f"=AVERAGE({VAN})", '#,##0;(#,##0);-'),
    ("Mediana", f"=MEDIAN({VAN})", '#,##0;(#,##0);-'),
    ("Desviación estándar", f"=STDEV({VAN})", '#,##0;(#,##0);-'),
    ("Coeficiente de variación", "=IFERROR(B9/B7,\"n/d\")", '0.00'),
    ("Mínimo", f"=MIN({VAN})", '#,##0;(#,##0);-'),
    ("Máximo", f"=MAX({VAN})", '#,##0;(#,##0);-'),
    ("Asimetría (skew)", f"=SKEW({VAN})", '0.000'),
    ("Curtosis (exceso)", f"=KURT({VAN})", '0.000'),
    ("IC 95% de la media – inferior", "=B7-1.96*B9/SQRT(B5)", '#,##0;(#,##0);-'),
    ("IC 95% de la media – superior", "=B7+1.96*B9/SQRT(B5)", '#,##0;(#,##0);-'),
    ("Prob. VAN < 0  (pérdida)", f'=COUNTIF({VAN},"<0")/COUNT({VAN})', '0.0%'),
    ("Prob. VAN ≥ 0  (proyecto viable)", "=1-B17", '0.0%'),
]
for i, (k, fml, fmt) in enumerate(stats):
    rr = 5 + i
    rs[f"A{rr}"] = k; rs[f"B{rr}"] = fml
    rs[f"A{rr}"].font = f_norm; rs[f"B{rr}"].font = f_bold
    rs[f"B{rr}"].number_format = fmt
    rs[f"A{rr}"].border = rs[f"B{rr}"].border = box
rs["A19"] = "Valor objetivo del VAN (edítalo)"; rs["B19"] = 20000
rs["B19"].font = f_in; rs["B19"].fill = fill_in; rs["B19"].number_format = '#,##0;(#,##0);-'
rs["A20"] = "Prob. VAN ≥ valor objetivo"; rs["B20"] = f'=COUNTIF({VAN},">="&B19)/COUNT({VAN})'
rs["B20"].number_format = '0.0%'; rs["B20"].font = f_bold
for a in ("A19", "A20"): rs[a].font = f_norm
for a in ("A19", "B19", "A20", "B20"): rs[a].border = box

# Percentiles
header(rs, 4, 4, ["Percentil", "VAN", "TIR"])
for i, p in enumerate([0.01, 0.05, 0.10, 0.25, 0.50, 0.75, 0.90, 0.95, 0.99]):
    rr = 5 + i
    rs[f"D{rr}"] = p; rs[f"D{rr}"].number_format = '0%'
    rs[f"E{rr}"] = f"=PERCENTILE({VAN},D{rr})"; rs[f"E{rr}"].number_format = '#,##0;(#,##0);-'
    rs[f"F{rr}"] = f'=IFERROR(PERCENTILE({TIR},D{rr}),"n/d")'; rs[f"F{rr}"].number_format = '0.0%'
    for a in ("D", "E", "F"):
        rs[f"{a}{rr}"].font = f_norm; rs[f"{a}{rr}"].border = box
rs["D15"] = "TIR media"; rs["F15"] = f'=IFERROR(AVERAGE({TIR}),"n/d")'; rs["F15"].number_format = '0.0%'
rs["D16"] = "TIR determinista"; rs["F16"] = "=Supuestos!B22"; rs["F16"].number_format = '0.0%'
rs["D17"] = "Corridas con TIR válida"; rs["F17"] = f"=COUNT({TIR})"
for a in ("D15", "D16", "D17"): rs[a].font = f_norm
for a in ("F15", "F16", "F17"): rs[a].font = f_bold

# Histograma
HB = 20
H0 = 24
rs[f"A{H0-1}"] = "Histograma del VAN"; rs[f"A{H0-1}"].font = f_bold
header(rs, H0, 1, ["Desde", "Hasta", "Centro", "Frecuencia", "% corridas", "% acumulado"])
for i in range(HB):
    rr = H0 + 1 + i
    rs[f"A{rr}"] = "=$B$11" if i == 0 else f"=B{rr-1}"
    rs[f"B{rr}"] = f"=A{rr}+($B$12-$B$11)/{HB}"
    rs[f"C{rr}"] = f'=TEXT((A{rr}+B{rr})/2,"#,##0")'
    if i < HB - 1:
        rs[f"D{rr}"] = f'=COUNTIFS({VAN},">="&A{rr},{VAN},"<"&B{rr})'
    else:
        rs[f"D{rr}"] = f'=COUNTIFS({VAN},">="&A{rr},{VAN},"<="&B{rr})'
    rs[f"E{rr}"] = f"=D{rr}/$B$5"
    rs[f"F{rr}"] = f"=SUM($E${H0+1}:E{rr})"
    for a, fmt in zip("ABCDEF", ['#,##0;(#,##0);-', '#,##0;(#,##0);-', '@', '#,##0', '0.0%', '0.0%']):
        rs[f"{a}{rr}"].number_format = fmt; rs[f"{a}{rr}"].font = f_norm; rs[f"{a}{rr}"].border = box
HL = H0 + HB

ch = BarChart(); ch.type = "col"; ch.title = "Distribución del VAN"
ch.y_axis.title = "Frecuencia"; ch.x_axis.title = "VAN (centro del intervalo)"
ch.add_data(Reference(rs, min_col=4, min_row=H0, max_row=HL), titles_from_data=True)
ch.set_categories(Reference(rs, min_col=3, min_row=H0 + 1, max_row=HL))
ch.legend = None; ch.gapWidth = 5; ch.height = 9; ch.width = 20
ch.x_axis.delete = False; ch.y_axis.delete = False
rs.add_chart(ch, "H4")

# Sensibilidad
S0 = HL + 3
rs[f"A{S0-1}"] = "Sensibilidad: correlación de cada variable con el VAN"; rs[f"A{S0-1}"].font = f_bold
header(rs, S0, 1, ["Variable", "Correlación", "|Corr.|"])
for i, v in enumerate(VARS):
    rr = S0 + 1 + i
    col = L(9 + i)
    rs[f"A{rr}"] = v[0]
    rs[f"B{rr}"] = f"=IFERROR(CORREL(Simulacion!${col}${R0}:${col}${R1},{VAN}),0)"
    rs[f"C{rr}"] = f"=ABS(B{rr})"
    rs[f"A{rr}"].font = f_norm
    rs[f"B{rr}"].number_format = '0.00'; rs[f"C{rr}"].number_format = '0.00'
    for a in "ABC": rs[f"{a}{rr}"].border = box
    rs[f"B{rr}"].font = f_bold; rs[f"C{rr}"].font = f_norm
SL = S0 + len(VARS)
rs[f"A{SL+1}"] = "Lectura: cuanto mayor es |correlación|, más influye esa variable en la incertidumbre del VAN."
rs[f"A{SL+1}"].font = f_sub

ch2 = BarChart(); ch2.type = "bar"; ch2.title = "Sensibilidad del VAN (correlación)"
ch2.add_data(Reference(rs, min_col=2, min_row=S0, max_row=SL), titles_from_data=True)
ch2.set_categories(Reference(rs, min_col=1, min_row=S0 + 1, max_row=SL))
ch2.legend = None; ch2.height = 9; ch2.width = 20
ch2.x_axis.delete = False; ch2.y_axis.delete = False
rs.add_chart(ch2, "H24")

for col, w in zip("ABCDEF", (36, 16, 14, 14, 12, 14)):
    rs.column_dimensions[col].width = w

wb.move_sheet("Instrucciones", offset=-wb.index(wb["Instrucciones"]))
wb.save("/Users/fernandev/Coding/myRiskSim/MiRiskSim.xlsx")
print("ok")
